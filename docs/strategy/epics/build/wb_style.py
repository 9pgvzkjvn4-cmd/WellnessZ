# -*- coding: utf-8 -*-
"""Shared styling helpers for the strategy workbooks (same look as the client's epic workbooks)."""
import math

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

NAVY = "FF1F3864"
F = "Arial"
thin = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
LABEL_FILL = PatternFill("solid", fgColor="FFDDEBF7")
METRIC_FILL = PatternFill("solid", fgColor="FFF2F2F2")
REVIEW_FILL = PatternFill("solid", fgColor="FFFFF2CC")
STATUS_FILLS = {"Done": "FFC6E0B4", "Partial": "FFFFE699", "Planned": "FFDDEBF7"}


def title_block(ws, title, subtitle, doc_line, last_col=None):
    ws["A1"] = title
    ws["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=F, size=10, color="FF595959")
    ws["A3"] = doc_line
    ws["A3"].font = Font(name=F, size=9, color="FF7F7F7F")
    ws.row_dimensions[1].height = 17.35
    ws.row_dimensions[2].height = 15.0
    ws.row_dimensions[3].height = 15.0
    if last_col:
        for r in (1, 2, 3):
            ws.merge_cells(f"A{r}:{last_col}{r}")


def header_row(ws, headers, widths, row=4):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=F, bold=True, size=10, color="FFFFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 31.5


def page_setup(ws, title_rows="4:4"):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = title_rows


def lines_needed(text, width):
    if text is None:
        return 1
    chars_per_line = max(1.0, width * 1.15)
    return sum(max(1, math.ceil(len(p) / chars_per_line)) for p in str(text).split("\n"))


def fit_row(ws, row, widths, minimum=15.0):
    n = max(lines_needed(ws.cell(row=row, column=i + 1).value, w) for i, w in enumerate(widths))
    ws.row_dimensions[row].height = max(minimum, round(n * 12.6 + 3, 2))


def body_cell(ws, row, col, value, center=False, wrap=True, review=False, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=F, size=10, color="FF000000", bold=bold)
    c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal="center" if center else "left")
    c.border = BORDER
    if review:
        c.fill = REVIEW_FILL
    elif fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def moscow_cf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Must"'], fill=PatternFill(bgColor="FFF8CBAD")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Should"'], fill=PatternFill(bgColor="FFFFE699")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Could"'], fill=PatternFill(bgColor="FFC6E0B4")))


def list_dv(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)
