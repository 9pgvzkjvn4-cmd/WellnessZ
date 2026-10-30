# -*- coding: utf-8 -*-
"""Builds one workbook per strategy epic (docs/strategy/epics/E2x_<Name>.xlsx) in the layout of the client's epic
workbooks (E09.xlsx): Epic Overview, High-Level Requirements, Feature List, Implementation Backlog, Legends.

Adapted from the E15 / E19 security epic builder (scratchpad sec/build/build_e15.py).
Usage: python3 build_epics.py <output directory>
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill

import strategy_data as D
from wb_style import (BORDER, F, LABEL_FILL, METRIC_FILL, NAVY, body_cell, fit_row, header_row, lines_needed, list_dv,
                      moscow_cf, page_setup, title_block)

OUT_DIR = sys.argv[1] if len(sys.argv) > 1 else ".."
DECISIONS = '"Accept,Modify,Reject,Defer"'

NOTES = [
    "How to review: yellow cells are for the reviewers named on the overview (C-suite sponsor, CIO, CMO / Nursing Head, Cyber GRC, DPO). Use the 'Decision' and 'Review comment' columns on each sheet; the dropdowns offer Accept / Modify / Reject / Defer.",
    "Story points are indicative relative-size estimates (Fibonacci-like scale) produced during strategy drafting; they must be re-estimated by the delivery teams. Team-sprints in the catalogue workbook use an indicative velocity of 30 points per team per two-week sprint.",
    "Horizons: H1 = Go-live readiness and close the gaps (months 0–6), H2 = Scale, intelligence and experience (months 6–18), H3 = Ecosystem and industry leadership (months 18–36). Priority on backlog items is P1 (critical) / P2 / P3.",
    "Nothing in this workbook exists in the product today unless the 'Evidence base' row says so; benefits and KPI targets are indicative and must be baselined by their business owners.",
    "Confidentiality rule used throughout: managers, HR, crew scheduling, vendors and analytics audiences see status or suppressed aggregates only – never diagnosis or clinical detail; all data stays in the client's jurisdiction.",
]

LEGENDS = [
    ("ALL", "All subsidiaries & departments (group-wide)"),
    ("FLT", "Flight Operations – pilots / flight crew"),
    ("CAB", "Cabin Services – cabin crew"),
    ("GND", "Ground Services – ramp, baggage, aircraft turnaround, GSE drivers"),
    ("CGO", "Cargo & Freight – warehouse, DGR/hazmat handling, cold chain, forklift"),
    ("CAT", "Catering – food production, food handlers, cold rooms, dish-wash, knives"),
    ("ENG", "Engineering & Maintenance – hangar, hazmat/chemicals, noise, heights, confined space"),
    ("SEC", "Security, Fire & Rescue, Medical Response"),
    ("H1", "Horizon 1 – Go-live readiness and close the gaps (months 0–6)"),
    ("H2", "Horizon 2 – Scale, intelligence and experience (months 6–18)"),
    ("H3", "Horizon 3 – Ecosystem and industry-leading OH platform (months 18–36)"),
    ("P1 / P2 / P3", "Backlog priority: P1 critical, P2 important, P3 desirable"),
    ("ISO45001", "ISO 45001:2018 – OH&S management systems"),
    ("ISO45003", "ISO 45003:2021 – psychological health & safety / psychosocial risk"),
    ("ISO27001", "ISO/IEC 27001:2022 – information security management, Annex A"),
    ("ISO27017", "ISO/IEC 27017 – information security controls for cloud services (tenant segregation)"),
    ("ISO27701", "ISO/IEC 27701 – privacy information management"),
    ("ISO22301", "ISO 22301 – business continuity management"),
    ("ISO20889", "ISO/IEC 20889 – privacy-enhancing data de-identification techniques"),
    ("ISO25012", "ISO/IEC 25012 – data quality model"),
    ("ISO42001", "ISO/IEC 42001 – AI management system"),
    ("ISO23894", "ISO/IEC 23894 – AI risk management"),
    ("NIST-AI-RMF", "NIST AI Risk Management Framework 1.0"),
    ("EU-AI-ACT", "Reg (EU) 2024/1689 (AI Act) – used as benchmark for risk classification, human oversight and transparency; local AI regulation applies"),
    ("DCB0129", "DCB0129 / DCB0160 – clinical risk management for health IT manufacture and deployment (UK standards, used as benchmark)"),
    ("SOC2", "AICPA Trust Services Criteria – CC (security), A (availability), C (confidentiality), PI, P"),
    ("SRE", "Site reliability engineering practice – service-level objectives, error budgets, burn-rate alerting"),
    ("CIS-K8S", "CIS Kubernetes / AKS benchmarks"),
    ("NIST-SSDF", "NIST SP 800-218 Secure Software Development Framework"),
    ("SLSA", "Supply-chain Levels for Software Artifacts – build provenance"),
    ("OWASP-MASVS", "OWASP Mobile Application Security Verification Standard v2"),
    ("WCAG22", "W3C Web Content Accessibility Guidelines 2.2, level AA"),
    ("EN301549", "EN 301 549 – accessibility requirements for ICT products and services"),
    ("PDF-UA", "ISO 14289 – accessible PDF"),
    ("HL7-FHIR", "HL7 FHIR R4 / R4B / R5 incl. subscriptions; SMART App Launch"),
    ("HL7-ODH", "HL7 FHIR Occupational Data for Health (ODH) IG"),
    ("HL7-V2", "HL7 version 2.5.1 messaging (ORU, ORM, ADT)"),
    ("IHE", "IHE IT Infrastructure profiles – MHD, XDS.b, PIXm, PDQm, ATNA, IUA"),
    ("SNOMED-ICD", "SNOMED CT / ICD-10 or ICD-11 / LOINC – clinical, diagnosis and lab coding"),
    ("EASA-MED", "EASA Part-MED (Reg (EU) 1178/2011 Annex IV) – pilot medical certificates, cabin crew medical reports"),
    ("ICAO-A1", "ICAO Annex 1 Ch. 6 – medical provisions for licensing"),
    ("ICAO-A6", "ICAO Annex 6 – operator responsibilities incl. fatigue management (FRMS); ICAO Doc 9966 FRMS manual"),
    ("ICAO-8984", "ICAO Doc 8984 – Manual of Civil Aviation Medicine"),
    ("IOSA", "IATA Operational Safety Audit – crew fitness and SMS standards"),
    ("ILO-C161", "ILO Convention C161 – occupational health services"),
    ("WHO-IHR", "WHO International Health Regulations – preparedness and reporting"),
    ("HIMSS", "HIMSS Digital Health Indicator (incl. interoperability, person-enabled health, predictive analytics)"),
    ("LOCAL-OHS", "Local labour law, national OH&S regulation and civil aviation authority medical requirements (e.g. UK CAA, UAE GCAA)"),
    ("LOCAL-PDP", "Local personal-data-protection law + GDPR Art. 9 (special-category health data) as benchmark"),
    ("Type: Story", "User-facing functionality expressed as a user story"),
    ("Type: Enabler", "Architectural / technical foundation work"),
    ("Type: Task", "Non-development work (content, configuration, governance, training, procurement)"),
    ("Type: Spike", "Time-boxed investigation to remove uncertainty"),
]


def build(e):
    wb = Workbook()
    wb.properties.title = f"{e['id']} – {e['name']}"
    wb.properties.subject = "WellnessZ OHCM – strategy epic workbook"
    wb.properties.creator = "WellnessZ Product & Architecture"
    hlrs, feats, bl = e["hlrs"], e["features"], e["backlog"]
    hlr_last, feat_last, bl_last = 4 + len(hlrs), 4 + len(feats), 4 + len(bl)

    # ---------------------------------------------------------------- Epic Overview
    ws = wb.active
    ws.title = "Epic Overview"
    ws.sheet_view.showGridLines = False
    title_block(ws, f"{e['id']} – {e['name']}",
                "Epic overview, scope, value, KPIs and summary metrics (review with the C-suite sponsor, CIO, CMO / Nursing Head and Cyber GRC)",
                D.DOC_LINE)
    ws.merge_cells("A1:D1")
    for col, w in (("A", 28.0), ("B", 95.0), ("C", 22.0)):
        ws.column_dimensions[col].width = w
    r = 5
    for label, value in e["overview"]:
        a = ws.cell(row=r, column=1, value=label)
        a.font = Font(name=F, bold=True, size=10, color=NAVY)
        a.fill = LABEL_FILL
        a.alignment = Alignment(wrap_text=True, vertical="top")
        a.border = BORDER
        b = ws.cell(row=r, column=2, value=value)
        b.font = Font(name=F, size=10, color="FF000000")
        b.alignment = Alignment(wrap_text=True, vertical="top")
        b.border = BORDER
        long = len(value) > 110 or "\n" in value
        if long:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            ws.cell(row=r, column=3).border = BORDER
        fit_row(ws, r, [28.0, 117.0 if long else 95.0])
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Summary metrics (formula-driven)").font = Font(name=F, bold=True, size=11, color=NAVY)
    r += 1
    HL, FL, BL = "'High-Level Requirements'", "'Feature List'", "'Implementation Backlog'"
    metrics = [
        ("Number of high-level requirements", f"=COUNTA({HL}!A5:A{hlr_last})"),
        ("Number of features", f"=COUNTA({FL}!A5:A{feat_last})"),
        ("Number of backlog items", f"=COUNTA({BL}!A5:A{bl_last})"),
        ("Total story points (indicative)", f"=SUM({BL}!G5:G{bl_last})"),
        ("Story points – H1 (0–6 months)", f'=SUMIF({BL}!H5:H{bl_last},"H1",{BL}!G5:G{bl_last})'),
        ("Story points – H2 (6–18 months)", f'=SUMIF({BL}!H5:H{bl_last},"H2",{BL}!G5:G{bl_last})'),
        ("Story points – H3 (18–36 months)", f'=SUMIF({BL}!H5:H{bl_last},"H3",{BL}!G5:G{bl_last})'),
        ("Indicative team-sprints (at 30 SP per sprint)", f"=ROUNDUP(SUM({BL}!G5:G{bl_last})/{D.VELOCITY},0)"),
        ("Must-have requirements", f'=COUNTIF({HL}!D5:D{hlr_last},"Must")'),
        ("Backlog items awaiting review decision", f"=COUNTBLANK({BL}!N5:N{bl_last})"),
    ]
    for label, formula in metrics:
        a = ws.cell(row=r, column=1, value=label)
        a.font = Font(name=F, bold=True, size=10, color="FF000000")
        a.fill = METRIC_FILL
        a.border = BORDER
        a.alignment = Alignment(wrap_text=True, vertical="top")
        b = ws.cell(row=r, column=2, value=formula)
        b.font = Font(name=F, size=10, color="FF000000")
        b.border = BORDER
        b.alignment = Alignment(horizontal="left")
        ws.row_dimensions[r].height = 15.0 if len(label) < 34 else 27.0
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Notes & assumptions").font = Font(name=F, bold=True, size=11, color=NAVY)
    r += 1
    for note in NOTES:
        c = ws.cell(row=r, column=1, value=note)
        c.font = Font(name=F, size=9, color="FF404040")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 30.0
        r += 1
    page_setup(ws)

    # ---------------------------------------------------------------- High-Level Requirements
    ws = wb.create_sheet("High-Level Requirements")
    title_block(ws, f"{e['id']} – High-Level Requirements",
                "What the solution must achieve for this epic (MoSCoW prioritised, traced to standards and to the evidence of increments 1–5)",
                D.DOC_LINE, "H")
    widths = [13.0, 70.0, 45.0, 12.0, 28.0, 34.0, 36.0, 12.0]
    header_row(ws, ["HLR ID", "Requirement", "Rationale", "Priority (MoSCoW)", "Standards / sources", "Applicable units",
                    "Review comment (sponsor / CMO / GRC)", "Decision"], widths)
    for i, (hid, req, rat, mos, std, units) in enumerate(hlrs):
        row = 5 + i
        for col, val, kw in ((1, hid, dict(center=True, wrap=False)), (2, req, {}), (3, rat, {}),
                             (4, mos, dict(center=True, wrap=False)), (5, std, {}), (6, units, {})):
            body_cell(ws, row, col, val, **kw)
        body_cell(ws, row, 7, None, review=True)
        body_cell(ws, row, 8, None, center=True, wrap=False, review=True)
        fit_row(ws, row, widths, 35.05)
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:H{hlr_last}"
    list_dv(ws, DECISIONS, f"H5:H{hlr_last}")
    moscow_cf(ws, f"D5:D{hlr_last}")
    page_setup(ws)

    # ---------------------------------------------------------------- Feature List
    ws = wb.create_sheet("Feature List")
    title_block(ws, f"{e['id']} – Low-Level Feature List",
                "Features that implement the high-level requirements, with roles, acceptance criteria and unit-specific notes",
                D.DOC_LINE, "J")
    widths = [12.0, 28.0, 60.0, 24.0, 12.0, 18.0, 48.0, 36.0, 30.0, 12.0]
    header_row(ws, ["Feature ID", "Feature", "Description", "Primary roles", "Priority (MoSCoW)", "Traces to HLR",
                    "Acceptance criteria (summary)", "Unit-specific notes", "Review comment", "Decision"], widths)
    for i, (fid, name, desc, roles, mos, hlr, ac, notes) in enumerate(feats):
        row = 5 + i
        for col, val, kw in ((1, fid, dict(center=True, wrap=False)), (2, name, {}), (3, desc, {}), (4, roles, {}),
                             (5, mos, dict(center=True, wrap=False)), (6, hlr, {}), (7, ac, {}), (8, notes, {})):
            body_cell(ws, row, col, val, **kw)
        body_cell(ws, row, 9, None, review=True)
        body_cell(ws, row, 10, None, center=True, wrap=False, review=True)
        fit_row(ws, row, widths, 35.05)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:J{feat_last}"
    list_dv(ws, DECISIONS, f"J5:J{feat_last}")
    moscow_cf(ws, f"E5:E{feat_last}")
    page_setup(ws)

    # ---------------------------------------------------------------- Implementation Backlog
    ws = wb.create_sheet("Implementation Backlog")
    title_block(ws, f"{e['id']} – Implementation Backlog (line by line)",
                "User stories, enablers, tasks and spikes with acceptance criteria, indicative estimates, horizon, priority, owner and dependencies",
                D.DOC_LINE, "O")
    widths = [14.0, 11.0, 26.0, 9.0, 60.0, 55.0, 10.0, 9.0, 9.0, 24.0, 14.0, 10.0, 13.0, 11.0, 30.0]
    header_row(ws, ["Backlog ID", "Feature ID", "Feature (lookup)", "Type", "Title / user story",
                    "Acceptance criteria / definition of done", "Story points (indicative)", "Horizon", "Priority",
                    "Owner / team", "Dependencies (epics)", "Sprint (to assign)", "Status", "Decision", "Review comment"], widths)
    for i, (fid, typ, title, ac, sp, hz, prio, owner, deps) in enumerate(bl):
        row = 5 + i
        body_cell(ws, row, 1, f"{e['id']}-BL-{i + 1:03d}", center=True, wrap=False)
        body_cell(ws, row, 2, fid, center=True, wrap=False)
        body_cell(ws, row, 3, f"=IFERROR(INDEX('Feature List'!$B$5:$B${feat_last},MATCH(B{row},'Feature List'!$A$5:$A${feat_last},0)),\"\")")
        body_cell(ws, row, 4, typ, center=True, wrap=False)
        body_cell(ws, row, 5, title)
        body_cell(ws, row, 6, ac)
        body_cell(ws, row, 7, sp, center=True, wrap=False)
        body_cell(ws, row, 8, hz, center=True, wrap=False)
        body_cell(ws, row, 9, prio, center=True, wrap=False)
        body_cell(ws, row, 10, owner)
        body_cell(ws, row, 11, deps)
        body_cell(ws, row, 12, None, wrap=False, review=True)
        body_cell(ws, row, 13, "Not started", center=True, wrap=False)
        body_cell(ws, row, 14, None, center=True, wrap=False, review=True)
        body_cell(ws, row, 15, None, review=True)
        feature_name = next(f[1] for f in feats if f[0] == fid)
        tmp = [f"{e['id']}-BL-000", fid, feature_name, typ, title, ac, str(sp), hz, prio, owner, deps]
        n = max(lines_needed(v, w) for v, w in zip(tmp, widths))
        ws.row_dimensions[row].height = max(23.85, round(n * 12.6 + 3, 2))
    tot = bl_last + 2
    for rr, label, formula in ((tot, "Total story points (indicative)", f"=SUM(G5:G{bl_last})"),
                               (tot + 1, "Items", f"=COUNTA(A5:A{bl_last})")):
        ws.cell(row=rr, column=5, value=label).font = Font(name=F, bold=True, size=10)
        c = ws.cell(row=rr, column=7, value=formula)
        c.font = Font(name=F, bold=True, size=10)
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:O{bl_last}"
    list_dv(ws, '"Not started,In analysis,Ready,In progress,Done,Deferred"', f"M5:M{bl_last}")
    list_dv(ws, DECISIONS, f"N5:N{bl_last}")
    list_dv(ws, '"H1,H2,H3"', f"H5:H{bl_last}")
    list_dv(ws, '"P1,P2,P3"', f"I5:I{bl_last}")
    list_dv(ws, '"Story,Enabler,Task,Spike"', f"D5:D{bl_last}")
    ws.conditional_formatting.add(f"I5:I{bl_last}", CellIsRule(operator="equal", formula=['"P1"'], fill=PatternFill(bgColor="FFF8CBAD")))
    page_setup(ws)

    # ---------------------------------------------------------------- Legends
    ws = wb.create_sheet("Legends")
    title_block(ws, "Legends – unit codes, standards, horizons, owners", "Reference for codes used in this workbook", D.DOC_LINE, "D")
    header_row(ws, ["Code", "Meaning"], [18.0, 120.0])
    for i, (code, meaning) in enumerate(LEGENDS):
        body_cell(ws, 5 + i, 1, code)
        body_cell(ws, 5 + i, 2, meaning)
        fit_row(ws, 5 + i, [18.0, 120.0])
    page_setup(ws)

    path = os.path.join(OUT_DIR, f"{e['id']}_{e['slug']}.xlsx")
    wb.save(path)
    return path, len(hlrs), len(feats), len(bl), sum(b[4] for b in bl)


if __name__ == "__main__":
    for ep in D.EPICS:
        p, h, f, b, sp = build(ep)
        print(f"{os.path.basename(p)}: HLR {h}, features {f}, backlog {b}, SP {sp}")
