# -*- coding: utf-8 -*-
"""Builds docs/strategy/epics/Epic-Catalogue-and-Roadmap.xlsx: Catalogue, Horizon roadmap, Dependencies, Capacity plan,
KPI tree. Usage: python3 build_catalogue.py <output directory>
"""
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import strategy_data as D
from wb_style import BORDER, F, NAVY, STATUS_FILLS, body_cell, fit_row, header_row, list_dv, page_setup, title_block

OUT_DIR = sys.argv[1] if len(sys.argv) > 1 else ".."
REFINE = {"H1": 1.5, "H2": 2.0, "H3": 2.5}   # cone-of-uncertainty allowance: first-cut backlogs grow on refinement
EVOLUTION_TS = 4      # team-sprints per quarter from Q3 for evolution of delivered modules (feedback, regulatory change)
QUARTERS = ["Q1 Nov–Jan 27", "Q2 Feb–Apr 27", "Q3 May–Jul 27", "Q4 Aug–Oct 27", "Q5 Nov–Jan 28", "Q6 Feb–Apr 28",
            "Q7 May–Jul 28", "Q8 Aug–Oct 28", "Q9 Nov–Jan 29", "Q10 Feb–Apr 29", "Q11 May–Jul 29", "Q12 Aug–Oct 29"]
HZ_OF_Q = ["H1"] * 2 + ["H2"] * 4 + ["H3"] * 6
HZ_FILL = {"H1": "FFF8CBAD", "H2": "FFFFE699", "H3": "FFC6E0B4"}
TEAMS_N = len(D.TEAMS)
TEAMS_BY_Q = [6, 6, 5, 5, 5, 5, 4, 4, 4, 4, 4, 4]   # stepped funding: six teams in H1, five in H2, four in H3
SPRINTS_PER_Q = 6.5
RESERVE = 0.35        # run / BAU, defects, sign-off support, security patching

# Platform completion (H1) – quarter spans
PLATFORM_Q = {"E13": (1, 2), "E15": (1, 3), "E16": (1, 2), "E17": (1, 2), "E18": (1, 2), "E19": (1, 4)}

# Non-epic H1 work items shown on the roadmap
H1_ITEMS = [
    ("GL-01", "Clinical, MRO and legal sign-offs (O11, O14, O19, O31; RTM-5 follow-up register)", 1, 2, "CMO, MRO, Legal, DPO"),
    ("GL-02", "Go-live configuration items (named clinicians O38, MFA enrolment O23 / O32, GitHub Environments O33)", 1, 1, "Platform operations, IAM"),
    ("GL-03", "Independent penetration test and retest (SEC-E01) – go-live condition G1", 2, 2, "CISO office"),
    ("GL-04", "Tier-0 two-region Azure landing zone and PROD go-live (16-week plan in the decision pack)", 1, 2, "Platform & SRE"),
    ("GL-05", "Retention classes approved and purge jobs enabled (D6, O12, O22, O27, O36)", 1, 2, "DPO, Legal"),
]

# Explicit dependencies: (from, to, type, reason)
DEPS = [
    ("E15", "E06", "Hard", "Absence feed connectors (HRIS, T&A, crew) – AbsenceFeedConnector is a no-op today"),
    ("E15", "E05", "Hard", "Rostering transport for the status-only fitness feed (O15)"),
    ("E15", "E10", "Hard", "Toxicology LIMS onboarding as bound integration client (O30)"),
    ("E15", "E08", "Hard", "Insurer FNOL, authority e-reporting and ERP connectors (D18, O24)"),
    ("E15", "E12", "Soft", "Campaign app and wearable aggregator (E12-F04 partial)"),
    ("E15", "E11", "Hard", "Public-health authority submission (default outbox + e-mail)"),
    ("E16", "E20", "Hard", "Customer-managed keys and envelope encryption before the PROD server and replica exist (SEC-E09, D5)"),
    ("E19", "E20", "Hard", "Signed-image admission and private control plane for the PROD clusters (SEC-E03, SEC-E04)"),
    ("E20", "Go-live", "Hard", "Two-region topology and first DR drill are Tier-0 go-live conditions"),
    ("E19", "Go-live", "Hard", "Go-live conditions G1–G6 (pen test, scan evidence, admission, demo-identity guard, CMK, patch SLAs)"),
    ("E13", "E22", "Hard", "Native mobile builds on the Capacitor shell and the employee app"),
    ("E19", "E22", "Hard", "Mobile hardening and MASVS assessment (SEC-E11)"),
    ("E18", "E21", "Soft", "Languages, templates and station list in master data"),
    ("E22", "E21", "Soft", "Platform accessibility APIs in the native app"),
    ("E15", "E23", "Hard", "Gateway, broker adapter and outbox are the transport of the hub"),
    ("E18", "E23", "Soft", "Terminology registry and value sets"),
    ("E05", "E24", "Hard", "Crew assessment and status-only roster feed"),
    ("E07", "E24", "Hard", "Crew medical certificates and validity"),
    ("E23", "E24", "Soft", "Document exchange and FHIR profiles for authority interfaces offering APIs"),
    ("E15", "E25", "Hard", "FRMS and rostering transport"),
    ("E26", "E25", "Soft", "Aggregated fatigue indicators via the data platform"),
    ("E17", "E26", "Hard", "De-identified view, KPI API and suppression are the starting point"),
    ("E15", "E26", "Hard", "CDC / broker feed into the analytics store"),
    ("E20", "E26", "Soft", "In-country analytics store in the approved region pair"),
    ("E26", "E27", "Hard", "No model without the governed, pseudonymised data platform"),
    ("E06", "E27", "Hard", "Early-intervention queue is where signals appear"),
    ("E27", "E28", "Soft", "Shared AI governance board and clinical safety process"),
    ("E03", "E28", "Hard", "AiDocumentationProvider extension point, consent purpose and audit"),
    ("E21", "E28", "Soft", "Arabic language support"),
    ("E18", "E29", "Hard", "Tenant configuration and content packs"),
    ("E20", "E29", "Hard", "Instances, overlays and upgrade orchestration"),
    ("E19", "E29", "Hard", "Extension signing and tenant isolation in the pen-test scope"),
    ("E26", "E30", "Hard", "Aggregated readiness data"),
    ("E27", "E30", "Soft", "Demand forecasts"),
    ("E14", "E30", "Hard", "Clinic capacity model"),
    ("E12", "E31", "Hard", "Campaign app, wearable and participation patterns"),
    ("E26", "E31", "Soft", "Spend and outcome analytics (ROI)"),
]

wb = Workbook()
wb.properties.title = "WellnessZ – Epic Catalogue and Roadmap"
wb.properties.subject = "WellnessZ OHCM – strategy beyond increment 5"
wb.properties.creator = "WellnessZ Product & Architecture"

new = {e["id"]: e for e in D.EPICS}

# ---------------------------------------------------------------------------------------------------------- Catalogue
ws = wb.active
ws.title = "Catalogue"
ws.sheet_view.showGridLines = False
title_block(ws, "Epic catalogue E01–E31", "All client epics, the security epic (renumbered E19) and the proposed strategy epics E20–E31 with status and evidence", D.DOC_LINE, "N")
widths = [8.0, 40.0, 14.0, 10.0, 22.0, 9.0, 11.0, 10.0, 9.0, 62.0, 22.0, 24.0, 11.0, 34.0]
header_row(ws, ["Epic", "Name", "Category", "Status", "Delivered by / horizon", "Items traced", "Implemented",
                "Sign-off pending", "Partial", "Evidence, gaps and goal", "Completion / delivery horizon", "Team",
                "Indicative SP", "Workbook"], widths)
rows = []
for (eid, name, status, by, notes, n, impl, so, part, hz) in D.EXISTING:
    cat = "Platform" if eid in ("E13", "E15", "E16", "E17", "E18") else ("Security" if eid == "E19" else "Business module")
    team = D.TEAM_OF[eid]
    sp = 256 if eid == "E19" else None
    wbk = "docs/security/epics/E15_Security_Hardening_and_Assurance.xlsx (to be renumbered E19)" if eid == "E19" else "Client workbook (scratchpad req*/)"
    if cat == "Platform":
        wbk = "Client catalogue (no workbook in repository)"
    rows.append((eid, name, cat, status, by, n, impl, so, part, notes, hz, team, sp, wbk))
for e in D.EPICS:
    sp = sum(b[4] for b in e["backlog"])
    rows.append((e["id"], e["name"], "Strategy (new)", "Planned", f"Proposed – {e['span']}", None, None, None, None,
                 e["one_line"], e["span"], e["team"], sp, f"docs/strategy/epics/{e['id']}_{e['slug']}.xlsx"))
rows.sort(key=lambda r: int(r[0][1:]))
for i, r in enumerate(rows):
    row = 5 + i
    for col, val in enumerate(r, start=1):
        body_cell(ws, row, col, val, center=col in (1, 4, 6, 7, 8, 9, 13), wrap=col not in (1,),
                  fill=STATUS_FILLS.get(val) if col == 4 else None)
    fit_row(ws, row, widths, 27.0)
last = 4 + len(rows)
r = last + 2
for label, formula in (("Done", f'=COUNTIF(D5:D{last},"Done")'), ("Partial", f'=COUNTIF(D5:D{last},"Partial")'),
                       ("Planned", f'=COUNTIF(D5:D{last},"Planned")'), ("Items traced (E01–E14)", f"=SUM(F5:F{last})"),
                       ("Implemented", f"=SUM(G5:G{last})"), ("Implemented – sign-off pending", f"=SUM(H5:H{last})"),
                       ("Partial", f"=SUM(I5:I{last})"), ("Indicative SP – E19 to E31", f"=SUM(M5:M{last})")):
    ws.cell(row=r, column=2, value=label).font = Font(name=F, bold=True, size=10)
    c = ws.cell(row=r, column=3, value=formula)
    c.font = Font(name=F, bold=True, size=10)
    r += 1
ws.cell(row=r + 1, column=1, value="Sources: release notes and requirements traceability matrices of increments 1–5 (per-epic counts; increment 5 distinguishes "
        "'implemented – configuration needs sign-off'); architecture document 01 Sections 1, 16–18; Security Enhancement Guideline. E05 also has one deferred training item. "
        "'Done' means functionally delivered and verified – go-live still depends on the sign-offs and connectors listed.").font = Font(name=F, size=9, color="FF404040")
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=14)
ws.cell(row=r + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[r + 1].height = 40
ws.freeze_panes = "C5"
ws.auto_filter.ref = f"A4:N{last}"
list_dv(ws, '"Done,Partial,Planned"', f"D5:D{last}")
page_setup(ws)

# ---------------------------------------------------------------------------------------------------- Horizon roadmap
ws = wb.create_sheet("Horizon roadmap")
ws.sheet_view.showGridLines = False
title_block(ws, "Horizon roadmap – 36 months from November 2026", "H1 go-live readiness (Q1–Q2) · H2 scale, intelligence and experience (Q3–Q6) · H3 ecosystem and industry leadership (Q7–Q12). Shaded = active.", D.DOC_LINE, "P")
widths = [9.0, 44.0, 24.0] + [8.5] * 12 + [30.0]
header_row(ws, ["ID", "Work item", "Team / owner"] + QUARTERS + ["Milestone / exit criterion"], widths)
for i, q in enumerate(QUARTERS):
    c = ws.cell(row=3, column=4 + i)
row = 5
MILESTONES = {
    "E13": "Store / MDM-ready app", "E15": "12 production connectors live", "E16": "Envelope encryption, anchoring, purge jobs",
    "E17": "Outstanding KPIs on KPI API", "E18": "Calendar, roster feed, maker-checker", "E19": "G1–G6 met; Later items closed",
    "E20": "Failover drill RTO ≤ 15 min proven", "E21": "WCAG 2.2 AA audit clean; Arabic everywhere", "E22": "Offline capture live at airside posts",
    "E23": "First live FHIR hospital + v2 lab", "E24": "Electronic submission to the first authority", "E25": "FRMS referral and roster restriction codes live",
    "E26": "All client KPIs from one semantic layer", "E27": "Silent pilot passes safety case → live pilot", "E28": "5-clinician pilot → staged roll-out",
    "E29": "Subsidiary onboarded in ≤ 4 weeks", "E30": "Readiness dashboard and scenarios", "E31": "Three vendors live with DPA",
}
items = []
for (gid, name, a, b, owner) in H1_ITEMS:
    items.append((gid, name, owner, a, b, "Go-live gate"))
for eid, (a, b) in PLATFORM_Q.items():
    nm = next(x[1] for x in D.EXISTING if x[0] == eid)
    team = next(p[3] for p in D.PLATFORM_COMPLETION if p[0] == eid)
    items.append((eid, f"{nm} – complete to production grade" if eid != "E19" else nm, team, a, b, MILESTONES[eid]))
for e in D.EPICS:
    items.append((e["id"], e["name"], e["team"], e["start_q"], e["end_q"], MILESTONES[e["id"]]))
for (iid, name, owner, a, b, ms) in items:
    body_cell(ws, row, 1, iid, center=True, wrap=False)
    body_cell(ws, row, 2, name)
    body_cell(ws, row, 3, owner)
    for q in range(1, 13):
        active = a <= q <= b
        body_cell(ws, row, 3 + q, "■" if active else None, center=True, fill=HZ_FILL[HZ_OF_Q[q - 1]] if active else None)
    body_cell(ws, row, 16, ms)
    fit_row(ws, row, widths, 27.0)
    row += 1
ws.freeze_panes = "D5"
page_setup(ws)

# ------------------------------------------------------------------------------------------------------- Dependencies
ws = wb.create_sheet("Dependencies")
ws.sheet_view.showGridLines = False
title_block(ws, "Dependencies between epics", "Predecessor → successor, hard (blocks start or go-live) or soft (reduces value or adds rework)", D.DOC_LINE, "F")
widths = [12.0, 12.0, 10.0, 80.0, 16.0, 16.0]
header_row(ws, ["Predecessor", "Successor", "Type", "Reason", "Predecessor status", "Successor horizon"], widths)
status = {x[0]: x[2] for x in D.EXISTING}
status.update({e["id"]: "Planned" for e in D.EPICS})
for i, (a, b, t, why) in enumerate(DEPS):
    row = 5 + i
    body_cell(ws, row, 1, a, center=True)
    body_cell(ws, row, 2, b, center=True)
    body_cell(ws, row, 3, t, center=True, fill="FFF8CBAD" if t == "Hard" else "FFFFE699")
    body_cell(ws, row, 4, why)
    st = status.get(a, "")
    body_cell(ws, row, 5, st, center=True, fill=STATUS_FILLS.get(st))
    body_cell(ws, row, 6, new[b]["span"] if b in new else ("H1" if b == "Go-live" else "Done – completion H1"), center=True)
    fit_row(ws, row, widths, 18.0)
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:F{4 + len(DEPS)}"
page_setup(ws)

# ------------------------------------------------------------------------------------------------------ Capacity plan
ws = wb.create_sheet("Capacity plan")
ws.sheet_view.showGridLines = False
title_block(ws, "Capacity plan – indicative team-sprints (not prices)",
            f"Velocity {D.VELOCITY} SP per team per two-week sprint; planning allowance ×1.5 / ×2.0 / ×2.5 for H1 / H2 / H3 epics (refinement, non-functional work, defects); "
            f"6 product teams in H1, 5 in H2, 4 in H3; {int(RESERVE * 100)} % of capacity reserved for run, defects and sign-off support", D.DOC_LINE, "S")
widths = [8.0, 38.0, 24.0, 9.0, 10.0, 11.0, 9.0] + [8.0] * 12
header_row(ws, ["Epic", "Name", "Team", "Horizon", "Indicative SP", "Planning team-sprints", "Quarters"] + QUARTERS, widths)
row = 5
plan_rows = []
for (eid, scope, ts, team, hz) in D.PLATFORM_COMPLETION:
    a, b = PLATFORM_Q[eid]
    plan_rows.append((eid, scope, team, hz, 256 if eid == "E19" else None, ts, a, b))
for e in D.EPICS:
    sp = sum(x[4] for x in e["backlog"])
    plan_rows.append((e["id"], e["name"], e["team"], e["span"], sp, math.ceil(sp * REFINE[e["horizon"]] / D.VELOCITY), e["start_q"], e["end_q"]))
plan_rows.append(("EVO", "Evolution of delivered modules E01–E14 (feedback, sign-off changes, regulatory change)", "T1 / T2", "H2–H3", None, EVOLUTION_TS * 10, 3, 12))
first = row
for (eid, name, team, hz, sp, ts, a, b) in plan_rows:
    body_cell(ws, row, 1, eid, center=True)
    body_cell(ws, row, 2, name)
    body_cell(ws, row, 3, team)
    body_cell(ws, row, 4, hz, center=True)
    body_cell(ws, row, 5, sp, center=True)
    if sp is not None and eid != "E19":
        body_cell(ws, row, 6, f"=ROUNDUP(E{row}*{REFINE[new[eid]['horizon']]}/{D.VELOCITY},0)", center=True, bold=True)
    else:
        body_cell(ws, row, 6, ts, center=True, bold=True)
    body_cell(ws, row, 7, f"{a}–{b}", center=True)
    span = b - a + 1
    for q in range(1, 13):
        col = 7 + q
        if a <= q <= b:
            body_cell(ws, row, col, f"=ROUND($F{row}/{span},1)", center=True, fill=HZ_FILL[HZ_OF_Q[q - 1]])
        else:
            body_cell(ws, row, col, None, center=True)
    fit_row(ws, row, widths, 27.0)
    row += 1
lastp = row - 1
row += 1
labels = [("Demand – roadmap team-sprints", lambda c: f"=SUM({c}{first}:{c}{lastp})"),
          ("Funded product teams", lambda c: TEAMS_BY_Q[ord(c) - ord("H")]),
          ("Capacity – gross team-sprints", lambda c: f"={c}{row - 1}*{SPRINTS_PER_Q}"),
          ("Capacity – available for roadmap", lambda c: f"=ROUND({c}{row - 1}*(1-{RESERVE}),1)"),
          ("Headroom for options (E32+) or client demand", lambda c: f"=ROUND({c}{row - 1}-{c}{rr['Demand – roadmap team-sprints']},1)"),
          ("Utilisation of available capacity", None)]
rr = {}
for label, fn in labels:
    ws.cell(row=row, column=2, value=label).font = Font(name=F, bold=True, size=10)
    rr[label] = row
    for q in range(1, 13):
        col = 7 + q
        L = get_column_letter(col)
        if fn:
            c = ws.cell(row=row, column=col, value=fn(L))
        else:
            c = ws.cell(row=row, column=col, value=f"=IFERROR({L}{rr['Demand – roadmap team-sprints']}/{L}{rr['Capacity – available for roadmap']},0)")
            c.number_format = "0%"
        c.font = Font(name=F, bold=True, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    row += 1
util_row = rr["Utilisation of available capacity"]
ws.conditional_formatting.add(f"H{util_row}:S{util_row}", CellIsRule(operator="greaterThan", formula=["1"], fill=PatternFill(bgColor="FFF8CBAD")))
ws.conditional_formatting.add(f"H{util_row}:S{util_row}", CellIsRule(operator="between", formula=["0.85", "1"], fill=PatternFill(bgColor="FFFFE699")))
row += 1
for t in D.TEAMS:
    ws.cell(row=row, column=2, value=f"{t[0]} {t[1]}").font = Font(name=F, bold=True, size=9)
    ws.cell(row=row, column=3, value=f"{t[2]} – {t[3]}").font = Font(name=F, size=9)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=12)
    row += 1
note = ("Reading the plan: utilisation above 100 % (red) means the quarter needs more capacity or a later start. H1 is the constrained horizon because the "
        "platform epics, E19 security and the Tier-0 build all land before go-live – decision S5 asks for a temporary integration capacity uplift in Q1–Q3. "
        "H2 and H3 epics carry a larger uncertainty allowance (×2.0 / ×2.5) because their backlogs are first cuts; the remaining headroom is deliberately "
        "unallocated for option epics (E32+), client demand and estimate overruns, and is reviewed at each horizon gate. Team-sprints are indicative and "
        "exclude business sign-off effort, external penetration testing, licences and hosting.")
c = ws.cell(row=row + 1, column=2, value=note)
c.font = Font(name=F, size=9, color="FF404040")
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=19)
ws.row_dimensions[row + 1].height = 42
ws.freeze_panes = "H5"
page_setup(ws)

# ------------------------------------------------------------------------------------------------------------ KPI tree
ws = wb.create_sheet("KPI tree")
ws.sheet_view.showGridLines = False
title_block(ws, "KPI tree – north star, outcomes and drivers", "Targets are indicative and taken from the client epic KPIs where available; baselines must be measured in the first 90 days of operation", D.DOC_LINE, "H")
widths = [12.0, 8.0, 8.0, 60.0, 40.0, 22.0, 20.0, 12.0]
header_row(ws, ["Level", "ID", "Parent", "KPI", "Target (indicative)", "Owner", "Epics", "Baseline"], widths)
fills = {"North star": "FF1F3864", "Outcome": "FFDDEBF7", "Driver": None}
for i, (lvl, kid, parent, kpi, tgt, owner, epics) in enumerate(D.KPI_TREE):
    row = 5 + i
    vals = [lvl, kid, parent, ("    " if lvl == "Driver" else "") + kpi, tgt, owner, epics, "to measure"]
    for col, v in enumerate(vals, start=1):
        c = body_cell(ws, row, col, v, center=col in (2, 3), fill=fills[lvl] if lvl != "North star" else None, bold=lvl != "Driver")
        if lvl == "North star":
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.font = Font(name=F, size=10, bold=True, color="FFFFFFFF")
    body_cell(ws, row, 8, "to measure", center=True, review=True)
    fit_row(ws, row, widths, 18.0)
ws.freeze_panes = "A5"
page_setup(ws)

out = os.path.join(OUT_DIR, "Epic-Catalogue-and-Roadmap.xlsx")
wb.save(out)
print("saved", out, "catalogue rows", len(rows), "roadmap items", len(items), "deps", len(DEPS), "capacity rows", len(plan_rows))
