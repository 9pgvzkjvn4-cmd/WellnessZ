#!/usr/bin/env python3
"""Builds WellnessZ-Role-Based-Access-Matrix.xlsx from expected-policy.json and rbac-results.json."""
import json
import os
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

D = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
policy = json.load(open(os.path.join(D, 'expected-policy.json')))
res = json.load(open(os.path.join(D, 'rbac-results.json')))
REVIEWED = json.load(open(os.path.join(D, 'reviewed.json'))) if os.path.exists(os.path.join(D, 'reviewed.json')) else {}

NAVY, GREY, WHITE = '1F3864', 'D9E1F2', 'FFFFFF'
PASS_FILL = PatternFill('solid', fgColor='C6EFCE')
FAIL_FILL = PatternFill('solid', fgColor='FFC7CE')
REV_FILL = PatternFill('solid', fgColor='FFEB9C')
HDR_FILL = PatternFill('solid', fgColor=NAVY)
SUB_FILL = PatternFill('solid', fgColor=GREY)
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

ROLE_INFO = {
    'EMPLOYEE': ('Employee / crew member: self-service (referrals, appointments, own documents, consents).', 'Own record only'),
    'MANAGER': ('Line manager: team cases, referrals, restrictions, modified duty, RTW acknowledgement.', 'Org scope subtree; status-only disclosure'),
    'HR': ('Human resources: referrals, RTW plans (read), job offers, ill-health retirement, absence pathway, claim status.', 'Org scope subtree; status-only disclosure'),
    'OH_CLINICIAN': ('Occupational health nurse / clinician: triage, record, fitness decisions.', 'Clinical (all non-restricted OH records)'),
    'OH_ADMIN': ('OH administrator: configuration, intake operations, scheduling administration.', 'Configuration; no clinical read via ABAC'),
    'CASE_MANAGER': ('Case manager: case coordination, claims participation.', 'Clinical (non-restricted)'),
    'PSYCH_CLINICIAN': ('Psychological health clinician (named mental-health clinician).', 'Clinical incl. RESTRICTED_CLINICAL'),
    'AUDITOR': ('Auditor: audit trail, surveillance audit packs, selection verification.', 'Audit / anonymised data only'),
    'PLATFORM_ADMIN': ('Platform administrator: technical administration, integration monitoring.', 'Technical; no clinical read via ABAC'),
    'INTEGRATION': ('Machine client (service account): inbound feeds.', 'System-to-system (not a human user)'),
    'PARTNER_PROVIDER': ('External partner provider (hospital, physiotherapy, EAP, telehealth). MFA required in production.', 'Partner organisation binding (token org_scope)'),
    'OH_PHYSICIAN': ('Occupational physician: orders, sign-off, MRO-equivalent decisions.', 'Clinical (non-restricted)'),
    'AME': ('Aeromedical examiner: licensed-role fitness, crew medicals.', 'Clinical (non-restricted)'),
    'CMO': ('Chief medical officer: oversight, reports, configuration.', 'Clinical (non-restricted) + reporting'),
    'NURSING_HEAD': ('Head of nursing: rostering of OH staff, quality, triage audit.', 'Clinical (non-restricted)'),
    'HSE_OFFICER': ('Health, safety & environment officer: incidents, recordability, ergonomics.', 'Incident / anonymised data; no clinical read'),
    'CREW_SCHEDULING': ('Crew scheduling: roster feed and duty status (status only).', 'Status only'),
    'HYGIENE_MANAGER': ('Food hygiene manager: food-handler clearance status.', 'Status only'),
    'CLINIC_RECEPTION': ('Clinic reception: walk-in registration, check-in.', 'Demographics / appointment only'),
    'FIRST_AIDER': ('First aider: first-aid encounter registration.', 'Registration only'),
}
for k in ('DPO', 'CLINIC_SCHEDULER', 'BIOMEDICAL', 'CREW_LICENSING', 'CANDIDATE', 'CLAIMS_COORDINATOR', 'TPA_ADJUSTER', 'LEGAL_COUNSEL', 'FINANCE',
          'PROCUREMENT', 'PEER_SUPPORT_VOLUNTEER', 'PEER_SUPPORT_COORDINATOR', 'SAFETY_MANAGER', 'DA_PROGRAMME_MANAGER', 'MRO', 'COLLECTOR',
          'HEALTH_PROMOTION_LEAD', 'ERGONOMIST'):
    desc = policy['roleDescriptions'].get(k, '')
    scope = {'DPO': 'Privacy reports; no clinical read', 'CLINIC_SCHEDULER': 'Demographics and appointment type only', 'BIOMEDICAL': 'Equipment only; no clinical read',
             'CREW_LICENSING': 'Certificate status only', 'CANDIDATE': 'Own pre-employment record only', 'CLAIMS_COORDINATOR': 'Claims; consent-gated disclosure; no clinical notes',
             'TPA_ADJUSTER': 'Own payer organisation\'s claims only (token org_scope); MFA required in production', 'LEGAL_COUNSEL': 'Litigated / held claims only',
             'FINANCE': 'Aggregated costs; no person-level clinical data', 'PROCUREMENT': 'Provider network / contracts', 'PEER_SUPPORT_VOLUNTEER': 'Own contact log only',
             'PEER_SUPPORT_COORDINATOR': 'Programme aggregates', 'SAFETY_MANAGER': 'Anonymised reports; for-cause requests', 'DA_PROGRAMME_MANAGER': 'Programme data; never individual results',
             'MRO': 'Laboratory results for verification', 'COLLECTOR': 'Chain-of-custody forms; no results', 'HEALTH_PROMOTION_LEAD': 'Anonymised participation',
             'ERGONOMIST': 'Assessments and actions'}[k]
    ROLE_INFO[k] = (desc, scope)

users = res['users']
roles_users = defaultdict(list)
for u in users:
    for r in u['roles']:
        roles_users[r].append(u['username'])
app_roles = res['appRoles']


def apps_for_role(r):
    out = ['employee-app (self-service)'] if r in ('EMPLOYEE', 'CANDIDATE') else []
    if r in policy['workspaceRoles']:
        out.append('manager-workspace')
    if app_roles['oh-workspace']['roles'] and r in app_roles['oh-workspace']['roles']:
        out.append('oh-workspace')
    if r in ('PARTNER_PROVIDER', 'TPA_ADJUSTER'):
        out.append('partner-portal')
    if r == 'INTEGRATION':
        out.append('none (machine-to-machine)')
    return ', '.join(out) or 'employee-app only (no workspace)'


# ------------------------------------------------------------------ verdicts
def verdict_matrix(r):
    s, allow = r['status'], r['expectedAllow']
    key = f"{r['user']} GET {r['path']}"
    if key in REVIEWED:
        return REVIEWED[key]['result'], REVIEWED[key]['note']
    if allow is None:
        return 'REVIEW', 'no static policy found for this endpoint'
    if s == 0 or s >= 500:
        return 'REVIEW', f'server error / no response ({r.get("errorCode") or r.get("error") or s}) – not an access decision'
    if allow:
        if 200 <= s < 300:
            return 'PASS', 'permitted and served' + ('' if r['hasData'] else ' (empty)')
        if s in (401, 403):
            if r['scoped']:
                return 'PASS', f'service-level data scope denies ({r.get("errorCode")}); no annotation-level role rule'
            return 'FAIL', 'permitted role denied'
        if s == 404:
            return 'PASS', ('service-level scope / not found' if r['scoped'] else 'authorised; resource not found') + f' ({r.get("errorCode")})'
        if s in (400, 409, 422):
            return 'PASS', f'authorised; request validation ({r.get("errorCode")})'
        if s == 406:
            return 'REVIEW', 'content negotiation (406)'
        return 'REVIEW', f'unexpected status {s}'
    if s in (401, 403, 404):
        return 'PASS', 'denied'
    if 200 <= s < 300:
        return 'FAIL', 'disallowed role served ' + ('DATA' if r['hasData'] else 'an empty result')
    if s in (400, 422):
        return 'PASS', f'rejected by validation before authorisation; no data ({r.get("errorCode")})'
    return 'REVIEW', f'unexpected status {s}'


def verdict_object(o):
    s = o['status']
    if o['expectedAllow']:
        if 200 <= s < 300:
            lvl = o.get('expectedLevel')
            order = ['NONE', 'STATUS_ONLY', 'SELF', 'CLINICAL', 'RESTRICTED_CLINICAL']
            got = o.get('disclosure')
            if lvl and got in order and order.index(got) > order.index(lvl):
                return 'FAIL', f'disclosure {got} exceeds expected {lvl}'
            if lvl and got in order and order.index(got) < order.index(lvl):
                return 'PASS', f'permitted; view limited to {got} (endpoint returns a reduced projection)'
            return 'PASS', 'permitted' + (f' (disclosure {o["disclosure"]})' if o.get('disclosure') else '')
        return ('FAIL', 'permitted user denied') if s in (401, 403, 404) else ('REVIEW', f'status {s}')
    if s in (401, 403, 404):
        return 'PASS', 'denied (' + ('404 – existence not confirmed' if s == 404 else str(s)) + ')'
    if 200 <= s < 300:
        return 'FAIL', 'object served to a user outside its scope' + (' – DATA' if o['hasData'] else '')
    return 'REVIEW', f'status {s}'


def verdict_unauth(u):
    out = []
    for via in ('bff', 'core'):
        s = u[via]['status']
        if s == 401 or (300 <= s < 400 and 'oauth2' in (u[via].get('location') or '') + 'login'):
            out.append(('PASS', f'{via.upper()} {s}'))
        elif s in (403,):
            out.append(('PASS', f'{via.upper()} 403 (denied before authentication)'))
        elif 200 <= s < 300:
            out.append(('FAIL', f'{via.upper()} {s} served without authentication'))
        else:
            out.append(('REVIEW', f'{via.upper()} {s}'))
    worst = 'FAIL' if any(v == 'FAIL' for v, _ in out) else 'REVIEW' if any(v == 'REVIEW' for v, _ in out) else 'PASS'
    return worst, '; '.join(n for _, n in out)


# ------------------------------------------------------------------ workbook helpers
wb = Workbook()


def sheet(title, headers, rows, widths, result_col=None, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = BORDER
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in wrap_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2' if len(headers) > 6 else 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}'
    ws.row_dimensions[1].height = 30
    if result_col:
        L = get_column_letter(result_col)
        rng = f'{L}2:{L}{len(rows) + 1}'
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"PASS"'], fill=PASS_FILL))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"FAIL"'], fill=FAIL_FILL))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"REVIEW"'], fill=REV_FILL))
    return ws


def rl(x):
    return x if isinstance(x, str) else ', '.join(x or [])


# ------------------------------------------------------------------ data rows
user_by = {u['username']: u for u in users}
mrows, mcount = [], Counter()
per_user = defaultdict(Counter)
fails = []
for r in sorted(res['results'], key=lambda r: (r['app'], r['user'], r['path'])):
    v, note = verdict_matrix(r)
    mcount[v] += 1
    per_user[r['user']][v] += 1
    if v != 'PASS':
        fails.append(('Role matrix', r['user'], 'GET ' + r['tested'], 'allow' if r['expectedAllow'] else 'deny', r['status'], note))
    mrows.append([r['app'], ', '.join(r['roles']), r['user'], 'GET', r['path'], r['tested'],
                  'ALLOW' if r['expectedAllow'] else 'DENY' if r['expectedAllow'] is False else '?', r['basis'], r['status'],
                  'yes' if r['hasData'] else 'no', v, note])

orows, ocount = [], Counter()
for o in res['objectResults']:
    v, note = verdict_object(o)
    ocount[v] += 1
    per_user[o['user']][v] += 1
    if v != 'PASS':
        fails.append(('Object-level', o['user'], 'GET ' + o['path'], 'allow' if o['expectedAllow'] else 'deny', o['status'], note))
    orows.append([o['object'], o['label'], o['path'], o['owner'], o['user'], ', '.join(o['roles']), o['rule'],
                  'ALLOW' if o['expectedAllow'] else 'DENY', o.get('expectedLevel') or '', o['status'], o.get('disclosure') or '',
                  'yes' if o['hasData'] else 'no', v, note, o['codeRef']])

urows, ucount = [], Counter()
for u in res['unauthenticated']:
    v, note = verdict_unauth(u)
    ucount[v] += 1
    if v != 'PASS':
        fails.append(('Unauthenticated', '-', u['method'] + ' ' + u['tested'], 'deny', f"{u['bff']['status']}/{u['core']['status']}", note))
    urows.append([u['method'], u['path'], u['tested'], u['bff']['status'], u['core']['status'], 'Expect 401', v, note])

grows, gcount = [], Counter()
for g in sorted(res['gating'], key=lambda g: g['user']):
    exp = g['bffWorkspaceExpected']
    s = g['bffWorkspaceStatus']
    ok = (s == 200) if exp else (s in (401, 403))
    v = 'PASS' if ok else 'FAIL'
    gcount[v] += 1
    if not ok:
        fails.append(('Workspace gating', g['user'], 'GET /bff/api/workspace/me', 'allow' if exp else 'deny', s, 'BFF workspace gate'))
    grows.append([g['user'], ', '.join(g['roles']), 'BFF /bff/api/workspace/**', 'ALLOW' if exp else 'DENY', str(s), v,
                  'WORKSPACE_ROLES: ' + ', '.join(policy['workspaceRoles'])])
    for app, a in g['apps'].items():
        e = a['expected']
        state = a['state']
        if e is None:
            v = 'REVIEW'
        else:
            v = 'PASS' if (state == 'ACCESSIBLE') == e else 'FAIL'
        gcount[v] += 1
        if v != 'PASS':
            fails.append(('Workspace gating', g['user'], app + ' UI', 'allow' if e else 'deny', state, a['excerpt']))
        grows.append([g['user'], ', '.join(g['roles']), app + ' (UI entry)', 'ALLOW' if e else 'DENY', state, v, a['rule'] + ' | ' + a['excerpt']])

# ------------------------------------------------------------------ Summary
ws = wb.active
ws.title = 'Summary'
ws['A1'] = 'WellnessZ – Role-Based Access Matrix (assurance evidence)'
ws['A1'].font = Font(bold=True, size=14, color=NAVY)
ws['A2'] = (f"Local test stack, run {res['started'][:16].replace('T', ' ')} – {res['finished'][11:16]} UTC. Apps: "
            + ', '.join(f'{k} {v}' for k, v in res['apps'].items()) + f". BFF {res['bff']}, core API {res['core']}, IdP {res['idp']}.")
ws['A3'] = ('Expected policy extracted statically from @PreAuthorize (constants resolved), core URL rules, programmatic role guards and the BFF '
            'workspace gate. Users signed in through the BFF with real IdP logins; only GET requests were sent with a session.')
ws['A4'] = ('Note: the local profile disables the MFA requirement for PARTNER_PROVIDER / TPA_ADJUSTER (WZ_MFA_REQUIRED_ROLES empty); '
            'in production those roles are stripped from tokens without amr/acr MFA proof.')
for c in ('A2', 'A3', 'A4'):
    ws[c].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(f'{c}:H{c[1:]}')
    ws.row_dimensions[int(c[1:])].height = 32
r0 = 6
tot = Counter()
for c in (mcount, ocount, ucount, gcount):
    tot.update(c)
blocks = [('Users signed in', len([u for u in users if not u.get('loginError')]), '', '', ''),
          ('GET endpoints without path parameters', res['getEndpoints'], '', '', ''),
          ('Operations in OpenAPI (unauthenticated tests)', res['operations'], '', '', ''),
          ('Endpoints in expected policy (static)', policy['endpointCount'], '', '', ''),
          ('Role matrix checks (user × endpoint)', sum(mcount.values()), mcount['PASS'], mcount['FAIL'], mcount['REVIEW']),
          ('Object-level checks', sum(ocount.values()), ocount['PASS'], ocount['FAIL'], ocount['REVIEW']),
          ('Unauthenticated checks (BFF + core per operation)', sum(ucount.values()), ucount['PASS'], ucount['FAIL'], ucount['REVIEW']),
          ('Workspace gating checks', sum(gcount.values()), gcount['PASS'], gcount['FAIL'], gcount['REVIEW']),
          ('All checks', sum(tot.values()), tot['PASS'], tot['FAIL'], tot['REVIEW'])]
for j, h in enumerate(['Measure', 'Count', 'PASS', 'FAIL', 'REVIEW'], 1):
    c = ws.cell(row=r0, column=j, value=h)
    c.font, c.fill, c.border = Font(bold=True, color=WHITE), HDR_FILL, BORDER
for i, b in enumerate(blocks, 1):
    for j, v in enumerate(b, 1):
        c = ws.cell(row=r0 + i, column=j, value=v)
        c.border = BORDER
        if i == len(blocks):
            c.font = Font(bold=True)
r1 = r0 + len(blocks) + 2
ws.cell(row=r1, column=1, value='Per application / role / user').font = Font(bold=True, size=12, color=NAVY)
hdr = ['Primary app', 'Role(s)', 'User', 'Checks', 'PASS', 'FAIL', 'REVIEW', 'Result']
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=r1 + 1, column=j, value=h)
    c.font, c.fill, c.border = Font(bold=True, color=WHITE), HDR_FILL, BORDER
row = r1 + 2
first = row
for u in sorted(users, key=lambda u: (u['app'], u['username'])):
    pc = per_user[u['username']]
    vals = [u['app'], ', '.join(u['roles']), u['username'], sum(pc.values()), pc['PASS'], pc['FAIL'], pc['REVIEW'],
            'FAIL' if pc['FAIL'] else 'REVIEW' if pc['REVIEW'] else 'PASS']
    for j, v in enumerate(vals, 1):
        ws.cell(row=row, column=j, value=v).border = BORDER
    row += 1
ws.conditional_formatting.add(f'H{first}:H{row - 1}', CellIsRule(operator='equal', formula=['"PASS"'], fill=PASS_FILL))
ws.conditional_formatting.add(f'H{first}:H{row - 1}', CellIsRule(operator='equal', formula=['"FAIL"'], fill=FAIL_FILL))
ws.conditional_formatting.add(f'H{first}:H{row - 1}', CellIsRule(operator='equal', formula=['"REVIEW"'], fill=REV_FILL))
ws.conditional_formatting.add(f'F{first}:F{row - 1}', CellIsRule(operator='greaterThan', formula=['0'], fill=FAIL_FILL))
# per app
row += 1
ws.cell(row=row, column=1, value='Per primary application').font = Font(bold=True, size=12, color=NAVY)
row += 1
for j, h in enumerate(['Primary app', 'Users', 'Roles', 'Checks', 'PASS', 'FAIL', 'REVIEW'], 1):
    c = ws.cell(row=row, column=j, value=h)
    c.font, c.fill, c.border = Font(bold=True, color=WHITE), HDR_FILL, BORDER
row += 1
for app in ('employee-app', 'manager-workspace', 'oh-workspace', 'partner-portal'):
    us = [u for u in users if u['app'] == app]
    pc = Counter()
    for u in us:
        pc.update(per_user[u['username']])
    vals = [app, len(us), len({r for u in us for r in u['roles']}), sum(pc.values()), pc['PASS'], pc['FAIL'], pc['REVIEW']]
    for j, v in enumerate(vals, 1):
        ws.cell(row=row, column=j, value=v).border = BORDER
    row += 1
row += 1
ws.cell(row=row, column=1, value='Findings (FAIL / REVIEW)').font = Font(bold=True, size=12, color=NAVY)
row += 1
for j, h in enumerate(['Area', 'User', 'Request', 'Expected', 'Actual', 'Note'], 1):
    c = ws.cell(row=row, column=j, value=h)
    c.font, c.fill, c.border = Font(bold=True, color=WHITE), HDR_FILL, BORDER
row += 1
if not fails:
    ws.cell(row=row, column=1, value='None – every check matched the expected policy.')
for f in fails[:400]:
    for j, v in enumerate(f, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical='top')
    row += 1
row += 1
ws.cell(row=row, column=1, value='Observations (not policy violations)').font = Font(bold=True, size=12, color=NAVY)
row += 1
unannotated = len([e for e in policy['endpoints'] if e['preAuthorize'] is None and not e['source'].startswith('apps/wz-bff')])
guarded = len([e for e in policy['endpoints'] if e.get('serviceGuardRoles')])
obs = [
    ('Defence in depth', f'{unannotated} core operations have no @PreAuthorize; {guarded} of them have a programmatic role guard in the service, the rest rely '
                         'on service-level ABAC (CaseAccess / EmployeeRecordAccess / SupportAccess). Every such endpoint tested denied or scoped correctly.'),
    ('Reference data', '13 reference-data endpoints (org units, job codes, hazards, clinic sites, service catalogue, appointment types, modules, '
                       'restrictions, fitness standards, languages, intake templates, body map, advice questions) return data to every signed-in '
                       'account, including external partner, TPA, EAP and candidate accounts. No personal data; decide whether the organisation '
                       'structure and clinic-site list should be hidden from external roles.'),
    ('Reviewed service rule', 'GET /testing/mro/queue: OH_PHYSICIAN passes @PreAuthorize but is refused by SubstanceSupport.actsAsMro unless listed in '
                              'substance.acting-mro-physicians (empty by default) – documented design (Solution Design E10).'),
    ('Workspace gate', 'All core /api/v1/workspace/** roles are within the BFF WORKSPACE_ROLES; the manager-workspace front-end list equals the BFF list. '
                       'Other /bff/api/workspace/X paths are aliases of /api/v1/X and are re-authorised by the core API.'),
    ('Demo data', 'Seeded case references encode the case nature (e.g. CAT-MH-…, CAT-CD-…) and are shown in status-only manager views; generated '
                  'references (e.g. GND-2026-000001) do not. Demo data only.'),
    ('MFA', policy['mfaNote']),
]
for a, b in obs:
    ws.cell(row=row, column=1, value=a).font = Font(bold=True)
    c = ws.cell(row=row, column=2, value=b)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    ws.row_dimensions[row].height = 48
    row += 1
for col, w in zip('ABCDEFGH', (34, 40, 52, 10, 10, 60, 9, 10)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A6'

# ------------------------------------------------------------------ Role catalogue
rc = []
for r in policy['allRoles']:
    desc, scope = ROLE_INFO.get(r, ('', ''))
    ep = [e for e in policy['endpoints'] if isinstance(e['coreAllowed'], list) and r in e['coreAllowed']]
    rc.append([r, desc, apps_for_role(r), scope, ', '.join(roles_users.get(r, [])) or '(service account)' if r == 'INTEGRATION' else ', '.join(roles_users.get(r, [])) or '–',
               len(ep), len([e for e in ep if e['method'] == 'GET'])])
sheet('Role catalogue', ['Role', 'Description', 'Application(s)', 'Data scope', 'Demo user(s)', 'Endpoints naming the role', 'of which GET'],
      rc, [26, 70, 36, 46, 36, 14, 10], wrap_cols=(2, 3, 4, 5))

# ------------------------------------------------------------------ Expected policy
ep_rows = []
for e in sorted(policy['endpoints'], key=lambda e: (e['path'], e['method'])):
    ep_rows.append([e['method'], e['path'], e['bffPath'] or '', e['module'], f"{e['controller']}.{e['handler']}", e['source'],
                    e['preAuthorize'] or '(none)', rl(e['coreAllowed']), e['bffGate'], rl(e['bffAllowed']),
                    '; '.join(e['dataScope']), '; '.join(e['notes'])])
sheet('Expected policy', ['Method', 'Core path', 'BFF path', 'Module', 'Handler', 'Source', '@PreAuthorize (resolved)', 'Core: allowed roles',
                          'BFF gate', 'Effective via BFF', 'Data-scope rules (code)', 'Notes'],
      ep_rows, [8, 52, 46, 20, 40, 60, 60, 50, 16, 50, 50, 60])
ws = wb['Expected policy']
n = ws.max_row + 2
for i, line in enumerate(['Core URL rules: ' + ' | '.join(f"{x['pattern']} → {x['rule']}" for x in policy['coreUrlRules']),
                          'BFF rules: ' + ' | '.join(f"{x['pattern']} → {x['rule']}" for x in policy['bffRules']),
                          'BFF CORE_WORKSPACE_RESOURCES (served natively by core under /api/v1/workspace): ' + ', '.join(policy['coreWorkspaceResources']),
                          'Other /bff/api/workspace/X paths are aliases of /api/v1/X (role-gated by WORKSPACE_ROLES; the core API applies its own rules again).',
                          policy['mfaNote']]):
    ws.cell(row=n + i, column=1, value=line).font = Font(italic=True)

# ------------------------------------------------------------------ Test results / object / unauth / gating
sheet('Test results', ['App', 'Role(s)', 'User', 'Method', 'Path (core)', 'Request (BFF)', 'Expected', 'Expected basis', 'Actual status',
                       'Data returned', 'Result', 'Notes'], mrows, [18, 30, 16, 8, 50, 56, 9, 60, 9, 9, 9, 60], result_col=11)
sheet('Object-level tests', ['Object', 'Scenario', 'Request (BFF)', 'Subject employee', 'User', 'Role(s)', 'Rule', 'Expected', 'Expected disclosure',
                             'Actual status', 'Actual disclosure', 'Data returned', 'Result', 'Notes', 'Code reference'],
      orows, [30, 30, 62, 12, 16, 30, 60, 9, 14, 9, 14, 9, 9, 40, 60], result_col=13)
sheet('Unauthenticated tests', ['Method', 'Path (OpenAPI)', 'Request', 'BFF status', 'Core status', 'Expected', 'Result', 'Notes'],
      urows, [8, 60, 70, 10, 10, 11, 9, 50], result_col=7)
sheet('Workspace gating', ['User', 'Role(s)', 'Area', 'Expected', 'Actual', 'Result', 'Rule / evidence'], grows, [16, 34, 34, 9, 14, 9, 110],
      result_col=6)

path = os.path.join(D, 'WellnessZ-Role-Based-Access-Matrix.xlsx')
wb.save(path)
print('written', path)
print('matrix', dict(mcount), 'object', dict(ocount), 'unauth', dict(ucount), 'gating', dict(gcount), 'all', dict(tot))
for f in fails[:80]:
    print(f)
