#!/usr/bin/env python3
"""Builds the WellnessZ demo guide workbook (WellnessZ-Demo-Guide.xlsx) from demo/scenarios.json and the demo data file.

usage: python3 demo/tools/build_demo_guide.py [out.xlsx]
"""
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

DEMO = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DEMO, 'dist', 'WellnessZ-Demo-Guide.xlsx')

APP_FILE = {'employee': 'employee-app.html', 'manager': 'manager-workspace.html', 'oh': 'oh-workspace.html',
            'partner': 'partner-portal.html'}
APP_NAME = {'employee': 'Employee & Crew App', 'manager': 'Manager & HR Workspace', 'oh': 'OH Clinical Workspace',
            'partner': 'Partner Portal'}

NAVY = '0D3B66'
HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
TITLE = Font(bold=True, size=16, color=NAVY)
SUB = Font(bold=True, size=12, color=NAVY)
WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Border(bottom=Side(style='thin', color='DBE3EE'))
APP_FILL = {'employee': 'E6F0FA', 'manager': 'EAF4E4', 'oh': 'FCEFE6', 'partner': 'E3F4F2'}


def seed():
    js = ("global.window={};require(process.argv[1]);"
          "process.stdout.write(JSON.stringify(window.WZ_DEMO_SEED))")
    return json.loads(subprocess.check_output(['node', '-e', js, os.path.join(DEMO, 'data', 'demo-data.js')]))


def table(ws, row, headers, rows, widths):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
    for r in rows:
        row += 1
        for c, v in enumerate(r, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.alignment, cell.border = WRAP, THIN
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    return row


def fmt_data(d):
    if not d:
        return ''
    return '\n'.join(f'{k}: {json.dumps(v) if not isinstance(v, str) else v}' for k, v in d.items())


def main():
    sc = json.load(open(os.path.join(DEMO, 'scenarios.json')))['scenarios']
    s = seed()
    users = {u['username']: u for u in s['users']}
    emps = {e['id']: e for e in s['employees']} if isinstance(s['employees'], list) else s['employees']
    wb = Workbook()

    ws = wb.active
    ws.title = 'How to run the demo'
    ws['A1'] = 'WellnessZ OHCM – Demo Pack: end-to-end demo guide'
    ws['A1'].font = TITLE
    ws['A2'] = f"Synthetic data only – demo data version {s['meta'].get('version', '')}. All people and organisations are fictional."
    lines = [
        ('Files', 'index.html (launcher), employee-app.html, manager-workspace.html, oh-workspace.html, partner-portal.html and '
                  'demo-data.js – keep all six files in the same folder.'),
        ('Browser', 'Google Chrome or Microsoft Edge (current version). Double-click index.html; no server, installation or '
                    'internet connection is needed.'),
        ('One shared session', 'All four apps share one demo session stored in the browser (localStorage key '
                               '"wellnessz-demo-session"). Changes made in one app appear in the others; open each app in its '
                               'own tab.'),
        ('Demo data file', 'demo-data.js is the synthetic starting data. "Export session" in the toolbar saves the current '
                           'session to wellnessz-demo-session.json; "Import" loads a saved session; "Reset" returns to the '
                           'starting data.'),
        ('Signing in', 'Use the user switcher in the demo toolbar at the top of each app. Only the users of that app are '
                       'listed (see sheet "Demo users").'),
        ('Guided mode', 'Pick a scenario in the toolbar to get a "next step" hint telling you which app and user to use. '
                        'Each scenario sheet lists the same steps with the expected result and a Pass/Fail column.'),
        ('Recommended order', 'Reset before each scenario for a clean run, or run S01–S10 in order in one session. '
                              'Approximate durations are on sheet "Scenarios".'),
        ('Rules shown in the demo', 'Managers and HR see status only; mental-health and drug-and-alcohol results are '
                                    'restricted classes (placeholder for other clinicians); partners see only their own '
                                    'organisation and nothing after consent withdrawal; small groups are suppressed in reports; '
                                    'every action is audited.'),
    ]
    r = 4
    for k, v in lines:
        ws.cell(row=r, column=1, value=k).font = SUB
        c = ws.cell(row=r, column=2, value=v)
        c.alignment = WRAP
        r += 1
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 120

    ws = wb.create_sheet('Scenarios')
    rows = []
    for i, x in enumerate(sc, 1):
        apps = sorted({st['app'] for st in x['steps']}, key=list(APP_FILE).index)
        rows.append([f'S{i:02d}', x['title'], ', '.join(APP_NAME[a] for a in apps),
                     '\n'.join(f"{p['user']} – {p['role']}" for p in x['personas']), len(x['steps']),
                     f"{max(4, len(x['steps']) * 1.5):.0f} min", f'S{i:02d} sheet'])
    table(ws, 1, ['ID', 'Scenario', 'Apps used', 'Personas (sign in as)', 'Steps', 'Approx. duration', 'Details'], rows,
          [7, 60, 40, 48, 8, 14, 12])
    ws.freeze_panes = 'A2'

    ws = wb.create_sheet('Demo users')
    rows = []
    for u in s['users']:
        e = emps.get(u.get('employeeId')) if u.get('employeeId') else None
        rows.append([u['username'], u['name'], ', '.join(u['roles']), ', '.join(APP_NAME[a] for a in u['apps']),
                     u.get('title', ''), u.get('employeeId') or '', u.get('orgScope') or '', u.get('partnerOrg') or '',
                     (e or {}).get('jobCode', '') if e else ''])
    table(ws, 1, ['Username', 'Name', 'Roles', 'App(s)', 'Title', 'Employee ID', 'Org scope', 'Partner organisation', 'Job code'],
          rows, [16, 24, 34, 34, 34, 12, 26, 26, 14])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for i, x in enumerate(sc, 1):
        ws = wb.create_sheet(f'S{i:02d}')
        ws['A1'] = f"S{i:02d} – {x['title']}"
        ws['A1'].font = TITLE
        ws['A2'] = 'Personas: ' + '; '.join(f"{p['user']} ({p['role']}, {APP_NAME[p['app']]})" for p in x['personas'])
        ws['A2'].alignment = WRAP
        rows = []
        for n, st in enumerate(x['steps'], 1):
            u = users.get(st['user'], {})
            rows.append([n, APP_NAME[st['app']], APP_FILE[st['app']], f"{st['user']} – {u.get('name', '')}",
                         st.get('action_label') or st['action'], st['expected'], fmt_data(st.get('data')), '', ''])
        last = table(ws, 4, ['Step', 'App', 'File', 'Sign in as', 'Action', 'Expected result', 'Test data used',
                             'Pass / Fail', 'Notes'], rows, [6, 22, 22, 26, 40, 60, 40, 11, 30])
        for rr in range(5, last + 1):
            app = x['steps'][rr - 5]['app']
            for cc in (2, 3):
                ws.cell(row=rr, column=cc).fill = PatternFill('solid', fgColor=APP_FILL[app])
        dv = DataValidation(type='list', formula1='"Pass,Fail,Blocked,Not run"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'H5:H{last}')
        ws.freeze_panes = 'A5'

    ws = wb.create_sheet('Rules checklist')
    rows = [
        ['Status only', 'Manager & HR Workspace', 'A manager or HR user never sees diagnosis, notes or results – only status, restrictions and dates', 'S01, S03, S04, S05'],
        ['Restricted class – mental health', 'OH Clinical Workspace', 'A general OH nurse sees a placeholder count, not the support case', 'S02, S10'],
        ['Restricted class – drug and alcohol', 'OH Clinical Workspace', 'Only the MRO sees the laboratory result; others see masked status', 'S04'],
        ['Employee choice', 'Employee & Crew App', 'Share fitness only: the manager sees exactly the fitness outcome', 'S02'],
        ['Partner isolation', 'Partner Portal', 'A partner sees only referrals for its own organisation; another partner cannot see them', 'S03, S10'],
        ['Consent', 'Employee & Crew App / Partner Portal', 'After consent withdrawal the partner loses access immediately', 'S03'],
        ['Small-cell suppression', 'Manager & HR Workspace', 'Group figures below 5 are suppressed', 'S07, S09'],
        ['Contact tracing privacy', 'Employee & Crew App', 'Contacts are notified without naming the case', 'S08'],
        ['Audit', 'OH Clinical Workspace (DPO)', 'Every action and every denied access is in the audit log', 'S02'],
    ]
    table(ws, 1, ['Rule', 'Where to see it', 'What to check', 'Scenarios'], rows, [34, 34, 80, 18])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print('written', OUT, len(sc), 'scenarios', sum(len(x['steps']) for x in sc), 'steps', len(s['users']), 'users')


if __name__ == '__main__':
    main()
